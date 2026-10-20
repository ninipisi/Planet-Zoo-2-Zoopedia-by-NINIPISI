"""Wandelt die Hilfsdatenbank (Excel) in die Daten der Webseite um.

Aufruf (im Projektordner):
  python tools/excel_to_json.py daten/Zoopedia_PZ2_Daten.xlsx            -> schreibt data/tiere.json, data/compat.json, data/meta.json
  python tools/excel_to_json.py daten/Zoopedia_PZ2_Daten.xlsx --pruefen  -> nur prüfen, nichts schreiben
Benötigt: pip install openpyxl

Gelesen werden die Blätter „Alle Tiere“ (eine Zeile je Tierart), „Inhalte“ (DLCs) und „Tiernamen“ (Übersetzungen).
Bei Fehlern wird nichts geschrieben – das Skript nennt Tierart und Spalte.
"""
import argparse
import collections
import datetime
import os
import re
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import zoopedia_schema as Z  # noqa: E402

errors, warnings = [], []
MISSING = collections.Counter()
N = Z.norm_head


def err(where, msg):
    errors.append(f"FEHLER   {where}: {msg}")


def warn(where, msg):
    warnings.append(f"Hinweis  {where}: {msg}")


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def prune(obj):
    """Leere Werte (None, leere Listen/Objekte) entfernen."""
    if isinstance(obj, dict):
        out = {}
        for k, v in obj.items():
            v = prune(v)
            if v is None or (isinstance(v, (dict, list)) and not v):
                continue
            out[k] = v
        return out
    return obj


def to_num(v, where, head):
    v = clean(v)
    if v is None or (isinstance(v, str) and v.startswith("=")):
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return v
    s = re.sub(r"(m²|m|%|°c|‰)$", "", str(v).lower().replace(",", ".").strip()).strip()
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        err(where, f"„{head}“: „{v}“ ist keine Zahl")
        return None


def to_bool(v, where, head):
    v = clean(v)
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in ("ja", "j", "x", "yes", "true", "1", "wahr"):
        return True
    if s in ("nein", "n", "no", "false", "0", "falsch", "-", "–"):
        return False
    err(where, f"„{head}“: „{v}“ – bitte Ja oder Nein")
    return None


RANGE_RE = re.compile(r"^\s*(\d+(?:[.,]\d+)?)\s*(?:[-–—]|bis)\s*(\d+(?:[.,]\d+)?)\s*$")


def to_range(v, where, head):
    """„1-6“ -> [1, 6]; „3“ -> [3, 3]"""
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return [v, v]
    m = RANGE_RE.match(str(v))
    if m:
        a, b = (float(x.replace(",", ".")) for x in m.groups())
        a, b = (int(x) if x.is_integer() else x for x in (a, b))
        return [min(a, b), max(a, b)]
    n = to_num(v, where, head)
    return [n, n] if n is not None else None


def to_time(v, where, head):
    """„3 Jahre“, „20 Monate“, „4 Wochen“, „10 Tage“ -> [Zahl, Einheit]; nackte Zahl = Jahre"""
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return [v, "years"]
    m = re.match(r"^\s*(\d+(?:[.,]\d+)?)\s*([a-zäöü]*)\.?\s*$", str(v).lower())
    if m:
        n = float(m.group(1).replace(",", "."))
        n = int(n) if n.is_integer() else n
        word = m.group(2) or "jahre"
        for uid, words in Z.TIME_UNITS:
            if word in words:
                return [n, uid]
    err(where, f"„{head}“: „{v}“ – bitte z. B. „3 Jahre“, „20 Monate“, „4 Wochen“ oder „10 Tage“")
    return None


def status_id(v, where):
    v = clean(v)
    if v is None:
        return None
    token = re.split(r"[\s(]", str(v).strip())[0].upper()
    for sid, abbr, name, iucn in Z.STATUSES:
        if token in (abbr.upper(), sid.upper(), iucn.upper()) or str(v).strip().lower() == name.lower():
            return sid
    err(where, f"Gefährdungslevel „{v}“ unbekannt")
    return None


def name_list(v):
    v = clean(v)
    if v is None or str(v).strip().lower() in ("keine", "-", "–", "nein"):
        return []
    return [x.strip() for x in re.split(r"[,;\n]", str(v)) if x.strip()]


# ---------- Überschriften erkennen ----------
def header_map(ws):
    """Spaltennummer -> (Gruppe, Unterspalte Zeile 3, Unterspalte Zeile 4), jeweils normalisiert.
    Verbundene Zellen gelten für alle Spalten, die sie überspannen; Gruppen laufen nach rechts weiter."""
    merged = {}
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= 4:
            v = ws.cell(mr.min_row, mr.min_col).value
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    merged[(r, c)] = v
    cell = lambda r, c: merged.get((r, c), ws.cell(r, c).value)
    cols, group, sub3 = {}, None, None
    for c in range(1, ws.max_column + 1):
        g = cell(1, c)
        if g is not None:
            group, sub3 = g, None
        r3 = cell(3, c)
        if r3 is not None:
            sub3 = r3
        elif (3, c) not in merged and cell(4, c) is None:
            sub3 = None
        cols[c] = (N(group), N(sub3 if r3 is None else r3), N(cell(4, c)))
    return cols


def find_cols(cols):
    """Ordnet die Excel-Spalten den Datenfeldern zu."""
    F = {}
    simple = {N(k): v for k, v in (("Tierart", "name"), ("Lat. Name", "sci"), ("Klasse", "class"), ("Ordnung", "order"),
                                     ("Familie", "family"), ("Gattung", "genus"), ("Bild (Dateiname)", "image"), ("Bild", "image"),
                                     ("Gefährdungslevel", "status"), ("Inhaltspaket", "pack"))}
    types = {N(lbl): tid for tid, lbl in Z.TYPES}
    roles = {N(lbl): rid for rid, lbl in Z.ROLES}
    conts = {N(lbl): cid for cid, lbl in Z.CONTINENTS}
    biomes = {N(lbl): bid for bid, lbl in Z.BIOMES}
    areas = {N(lbl): aid for aid, lbl in Z.AREAS}
    areas[N("Kletterfläche")] = "climb"
    life = {N(lbl): lid for lid, lbl in Z.LIFE}
    for c, (g, r3, r4) in cols.items():
        key = None
        if g in simple and r3 in ("", g) and r4 in ("", g):
            key = simple[g]
        elif g.startswith("gehegeart") and r3 in types:
            key = ("type", types[r3])
        elif g.startswith("ökologischerolle") and (r4 or r3) in roles:
            key = ("role", roles[r4 or r3])
        elif g.startswith("kontinente") and r3 in conts:
            key = ("continent", conts[r3])
        elif g == "biome" and r3 in biomes:
            key = ("biome", biomes[r3])
        elif g == "temperatur" and r3 in ("von", "bis"):
            key = "tempMin" if r3 == "von" else "tempMax"
        elif g.startswith("luftfeuchtigkeit") and r3 in ("von", "bis"):
            key = "humMin" if r3 == "von" else "humMax"
        elif g.startswith("salzgehalt"):
            key = "salinity"
        elif g.startswith("minimalegehegevoraussetzungen") and r3 in areas:
            if "je1" in g:
                key = ("area1", areas[r3])
            elif "je2" in g:
                key = ("area2", areas[r3])
        elif g.startswith("zaun"):
            if r3 == "kletterschutz":
                key = "climbProof"
            elif r3 == "zaunstärke":
                key = "strength"
            elif r3 == "zaunhöhe":
                key = "height"
            elif r3.startswith("wassergraben"):
                key = {"erlaubt": "moat", "breite": "moatW", "tiefe": "moatD"}.get(r4)
        elif g.startswith("sozialverhalten"):
            if r3 == "gemischtmin":
                key = "grpMin"
            elif r3 == "gemischtmax":
                key = "grpMax"
            elif r3.startswith("gemischtverhältnis"):
                key = {"m": "maxM", "w": "maxF"}.get(r4)
            elif r3 == "nurm":
                key = "onlyM"
            elif r3 == "nurw":
                key = "onlyF"
            elif "durchgehen" in r3 or "begehbar" in r3:
                key = "walk"
        elif g.startswith("lebenszyklus") and r3 in life:
            key = ("life", life[r3])
        elif g == "vivarium" and r3 == "bauformen":
            key = "enclosures"
        elif g.startswith("farbmorphen"):
            if r3.startswith("farbmorph"):
                key = "morphText"
            elif r3.startswith("beispielbild"):
                key = "morphImg"
        elif g.startswith("zusammenleben"):
            if r3.startswith("erfolgreich"):
                key = "thriving"
            elif r3.startswith("neutral"):
                key = "neutral"
        if key is not None:
            if key in F:
                err(Z.MAIN_SHEET, f"Spalte „{key}“ kommt doppelt vor")
            F[key] = c
    for need, head in (("name", "Tierart"), ("status", "Gefährdungslevel"), ("pack", "Inhaltspaket"),
                       ("tempMin", "Temperatur von"), ("tempMax", "Temperatur bis")):
        if need not in F:
            err(Z.MAIN_SHEET, f"Spalte „{head}“ nicht gefunden – wurde die Überschrift umbenannt?")
    for tid, lbl in Z.TYPES:
        if ("type", tid) not in F:
            err(Z.MAIN_SHEET, f"Spalte „Gehegeart – {lbl}“ nicht gefunden")
    return F


# ---------- Inhalte, Namen ----------
def read_dlcs(wb):
    ws = wb["Inhalte"]
    dlcs, by_name = [], {}
    for r, row in enumerate(ws.iter_rows(min_row=2, max_col=5, values_only=True), 2):
        did, de, en, active, url = (clean(x) for x in (tuple(row) + (None,) * 5)[:5])
        if not did and not de:
            continue
        did = str(did) if did else Z.slug(de)
        entry = {"id": did, "default": to_bool(active, f"Inhalte Zeile {r}", "Standardmäßig aktiv") is True}
        if url:
            url = str(url).strip()
            if not url.lower().startswith(("https://", "http://")):
                err(f"Inhalte Zeile {r}", f"Steam-Link „{url}“ muss mit https:// beginnen")
            else:
                entry["url"] = url
        if did not in ("base", "deluxe"):  # Basisspiel/Deluxe kommen aus den Übersetzungen
            entry["names"] = {"de": de or did, "en": en or de or did}
        if any(d["id"] == did for d in dlcs):
            err(f"Inhalte Zeile {r}", f"ID „{did}“ doppelt")
        dlcs.append(entry)
        for key in (did, de, en):
            if key:
                by_name[str(key).lower()] = did
    by_name.setdefault("basisspiel", "base")
    if not dlcs:
        err("Inhalte", "keine Inhalte eingetragen")
    return dlcs, by_name


def read_names(wb):
    """Blatt Tiernamen -> { deutscher Name (klein): { sprache: name } }"""
    out = {}
    if Z.NAMES_SHEET not in wb.sheetnames:
        return out
    ws = wb[Z.NAMES_SHEET]
    header = [clean(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    cols = {}
    for i, hname in enumerate(header):
        for code, label in Z.NAME_LANGS:
            if hname == Z.name_head(code, label) or hname == code:
                cols[i] = code
    for row in ws.iter_rows(min_row=2, values_only=True):
        de = clean(row[0]) if row else None
        if not de:
            continue
        names = {code: str(clean(row[i])) for i, code in cols.items() if i < len(row) and clean(row[i])}
        if names:
            out[re.sub(r"\s+", " ", str(de)).lower()] = names
    return out


def read_morph_texts(wb):
    """Blatt „Farbmorphen (Sprachen)“ -> { (deutscher Tiername klein, Nr): { sprache: "Name: Beschreibung" } }
    Spalten: Tierart | Nr | Deutsch (de) | English (en) | … – Nr = Reihenfolge wie in „Alle Tiere“ (1, 2, …)."""
    out = {}
    if Z.MORPHS_SHEET not in wb.sheetnames:
        return out
    ws = wb[Z.MORPHS_SHEET]
    header = [clean(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    cols = {}
    for i, hname in enumerate(header):
        for code, label in Z.NAME_LANGS + [("de", "Deutsch")]:
            if hname == Z.name_head(code, label) or hname == code:
                cols[i] = code
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not clean(row[0]) or clean(row[1]) is None:
            continue
        try:
            nr = int(float(clean(row[1])))
        except ValueError:
            continue
        texts = {code: str(clean(row[i])) for i, code in cols.items() if i < len(row) and clean(row[i])}
        out[(re.sub(r"\s+", " ", str(clean(row[0]))).lower(), nr)] = texts
    return out


def read_info(wb, key):
    """Blatt „Info“: Wert der Zeile mit der Angabe „key“ (z. B. Version, GoatCounter-Code)."""
    if "Info" not in wb.sheetnames:
        return None
    for row in wb["Info"].iter_rows(min_row=1, max_col=2, values_only=True):
        if row and str(clean(row[0]) or "").lower() == key.lower() and clean(row[1]) is not None:
            return str(clean(row[1]))
    return None


def read_updated(wb):
    """Blatt „Info“, Zeile „Letztes Update“ -> Datum oben links neben der Version (leer = kein Datum)."""
    v = None
    if "Info" in wb.sheetnames:
        for row in wb["Info"].iter_rows(min_row=1, max_col=2, values_only=True):
            if row and str(clean(row[0]) or "").lower() == "letztes update":
                v = row[1]
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        m = re.match(r"^\s*(\d{1,2})\.(\d{1,2})\.(\d{4})\s*$", v)
        if m:
            return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
        warn("Info", f"„Letztes Update“: „{v}“ ist kein Datum (TT.MM.JJJJ) – kein Datum angezeigt")
    return None


def read_version(wb):
    return read_info(wb, "Version")


# ---------- Tiere ----------
ENCLOSURE_LABELS = ["Box", "Zylinder", "Box (groß)", "Zylinder (groß)", "Box (groß, abgeschrägt)", "Begehbar"]
ENCLOSURES = dict(zip([x.lower() for x in ENCLOSURE_LABELS], ["box", "cylinder", "boxLarge", "cylinderLarge", "boxLargeAngled", "walkthrough"]))
AREA_HEAD = dict(Z.AREAS)


def read_animals(wb, dlc_by_name, image_dir):
    ws = wb[Z.MAIN_SHEET]
    F = find_cols(header_map(ws))
    if errors:
        return []
    animals = []
    for r in range(Z.FIRST_ROW, ws.max_row + 1):
        def val(k):
            return ws.cell(r, F[k]).value if k in F else None

        name = clean(val("name"))
        if not name:
            continue
        name = re.sub(r"\s+", " ", str(name))
        where = f"{name} (Zeile {r})"
        a = {"id": Z.slug(name), "name": name, "sci": clean(val("sci"))}

        pack = clean(val("pack")) or "Basisspiel"
        m = re.match(r"^\s*jubil(?:ä|ae)um\s*(\d{4})?\s*$", str(pack), re.I)
        if m:
            a["dlc"] = "base"  # Jubiläumstiere sind kostenlos – sie gehören zum Basisspiel
            a["gift"] = True
            if m.group(1):
                a["giftYear"] = int(m.group(1))
        else:
            did = dlc_by_name.get(str(pack).lower())
            if not did:
                err(where, f"Inhaltspaket „{pack}“ steht nicht im Blatt „Inhalte“")
            a["dlc"] = did or "base"
        a["status"] = status_id(val("status"), where)
        a["taxonomy"] = {k: clean(val(k)) for k in ("class", "order", "family", "genus")}

        img = clean(val("image"))
        if isinstance(img, str) and img.startswith("="):
            img = name + ".jpg"  # Formel =A5&".jpg"
        if img:
            stem = Z.slug(os.path.splitext(str(img))[0])
            if os.path.exists(os.path.join(image_dir, stem + ".webp")):
                a["image"] = stem
            else:
                MISSING["Tierbild"] += 1

        def x(k):
            return k in F and bool(clean(ws.cell(r, F[k]).value))

        a["types"] = [tid for tid, _ in Z.TYPES if x(("type", tid))]
        if not a["types"]:
            err(where, "keine Gehegeart angekreuzt")
        a["roles"] = [rid for rid, _ in Z.ROLES if x(("role", rid))]
        a["continents"] = [cid for cid, _ in Z.CONTINENTS if x(("continent", cid))]
        a["biomes"] = [bid for bid, _ in Z.BIOMES if x(("biome", bid))]

        def num(k, head):
            return to_num(val(k), where, head)

        a["tempMin"], a["tempMax"] = num("tempMin", "Temperatur von"), num("tempMax", "Temperatur bis")
        if a["tempMin"] is None:
            MISSING["Temperatur"] += 1
        h1, h2 = num("humMin", "Luftfeuchtigkeit von"), num("humMax", "Luftfeuchtigkeit bis")
        if h1 is not None or h2 is not None:
            a["humidity"] = [h1 if h1 is not None else h2, h2 if h2 is not None else h1]
        a["salinity"] = num("salinity", "Salzgehalt")

        for n, cnt in (("area1", "1 Tier"), ("area2", "2 Tiere")):
            ar = {aid: num((n, aid), f"{AREA_HEAD[aid]} ({cnt})") for aid, _ in Z.AREAS}
            ar = {k: v for k, v in ar.items() if v is not None}
            if ar:
                a[n] = ar
        if "area1" not in a and any(t in a["types"] for t in ("land", "flying")):
            MISSING["Flächen"] += 1

        moat_ok = to_bool(val("moat"), where, "Wassergraben erlaubt")
        mw, md = num("moatW", "Wassergraben Breite"), num("moatD", "Wassergraben Tiefe")
        fence = {"climbProof": to_bool(val("climbProof"), where, "Kletterschutz"),
                 "strength": num("strength", "Zaunstärke") or None, "height": num("height", "Zaunhöhe") or None}
        if moat_ok and mw:
            fence["moat"] = {"width": mw, "depth": md or 0}
        a["fence"] = fence

        a["social"] = {"min": num("grpMin", "Gemischt MIN"), "max": num("grpMax", "Gemischt MAX"),
                       "maxMales": num("maxM", "Verhältnis ♂"), "maxFemales": num("maxF", "Verhältnis ♀"),
                       "maleOnly": to_range(val("onlyM"), where, "Nur ♂"), "femaleOnly": to_range(val("onlyF"), where, "Nur ♀")}
        a["walkthrough"] = to_bool(val("walk"), where, "Gäste können durchgehen")

        lc = {}
        for lid, head in Z.LIFE:
            v = val(("life", lid))
            lc[lid] = to_range(v, where, head) if lid == "litter" else to_time(v, where, head)
        a["lifecycle"] = lc

        enc = [x.strip() for x in str(clean(val("enclosures")) or "").split(";") if x.strip()]
        a["enclosures"] = []
        for e in enc:
            eid = ENCLOSURES.get(e.lower())
            if eid:
                a["enclosures"].append(eid)
            else:
                err(where, f"Bauform „{e}“ unbekannt (erlaubt: {', '.join(ENCLOSURE_LABELS)})")
        texts = [x.strip() for x in str(clean(val("morphText")) or "").split(";") if x.strip()]
        imgs = [x.strip() for x in str(clean(val("morphImg")) or "").split(";") if x.strip()]
        if imgs and len(imgs) != len(texts):
            warn(where, f"Farbmorphen: {len(texts)} Beschreibung(en), aber {len(imgs)} Bild(er) – bitte gleiche Anzahl und Reihenfolge")
        a["morphs"] = []
        for i, tx in enumerate(texts):
            m = {"text": tx}
            if i < len(imgs):
                url = imgs[i]
                local = os.path.join("assets", "farbmorphen", os.path.basename(url.split("?")[0]))
                m["img"] = local.replace(os.sep, "/") if os.path.exists(os.path.join(Z.ROOT, local)) else url
            a["morphs"].append(m)

        a["_thriving"] = name_list(val("thriving"))
        a["_neutral"] = name_list(val("neutral"))
        animals.append(a)
    return animals


def build_compat(animals):
    """'benefit': a profitiert von b (aus „Erfolgreiche Beziehungen“ von a); 'neutral': verträglich, gilt in beide Richtungen."""
    by_name = {a["name"].lower(): a["id"] for a in animals}
    by_name.update({a["id"]: a["id"] for a in animals})
    pairs, seen = [], set()

    def add(a, b, typ):
        k = (a, b, typ) if typ == "benefit" else (min(a, b), max(a, b), typ)
        if k not in seen:
            seen.add(k)
            pairs.append({"a": k[0], "b": k[1], "type": typ})

    for a in animals:
        for key, typ in (("_thriving", "benefit"), ("_neutral", "neutral")):
            for nm in a.pop(key):
                other = by_name.get(nm.lower()) or by_name.get(Z.slug(nm))
                if not other:
                    err(a["name"], f"Beziehung zu „{nm}“ – diese Tierart steht nicht in der Liste (Schreibweise prüfen)")
                elif other != a["id"]:
                    add(a["id"], other, typ)
    return pairs


def main():
    ap = argparse.ArgumentParser(description="Hilfsdatenbank (Excel) -> data/*.json")
    ap.add_argument("excel")
    ap.add_argument("--pruefen", action="store_true", help="nur prüfen, nichts schreiben")
    args = ap.parse_args()

    meta = Z.load_json("data/meta.json")
    wb = load_workbook(args.excel)  # Formeln (Bild-Spalte) werden selbst ausgewertet
    if Z.MAIN_SHEET not in wb.sheetnames:
        print(f"FEHLER: Blatt „{Z.MAIN_SHEET}“ fehlt in der Excel.")
        sys.exit(1)
    dlcs, dlc_by_name = read_dlcs(wb)
    animals = read_animals(wb, dlc_by_name, os.path.join(Z.ROOT, "assets", "tiere"))

    for i, n in collections.Counter(a["id"] for a in animals).items():
        if n > 1:
            err(Z.MAIN_SHEET, f"Tierart „{i}“ steht {n}× in der Liste")
    pairs = build_compat(animals) if not errors else []

    names = read_names(wb)
    untranslated = 0
    for a in animals:
        n = names.get(a["name"].lower())
        if n:
            a["names"] = n
        else:
            untranslated += 1
    mt = read_morph_texts(wb)
    open_morphs = 0
    for a in animals:
        for i, m in enumerate(a.get("morphs") or []):
            tx = {k: v for k, v in mt.get((a["name"].lower(), i + 1), {}).items() if k != "de"}
            if tx:
                m["texts"] = tx
            if len(tx) < len(Z.NAME_LANGS) - 1:  # es-419 darf leer sein (nutzt es)
                open_morphs += 1
    if open_morphs:
        warn(Z.MORPHS_SHEET, f"{open_morphs} Farbmorph(e) nicht in allen Sprachen übersetzt – dort wird Englisch bzw. Deutsch angezeigt")
    if untranslated:
        warn(Z.NAMES_SHEET, f"{untranslated} Tierart(en) ohne Übersetzung – dort wird überall der deutsche Name angezeigt")
    if MISSING:
        warnings.append("Hinweis  noch offen – " + ", ".join(f"{k} bei {n} Tierart(en)" for k, n in MISSING.items()))
    for line in warnings + errors:
        print(line)
    cnt = collections.Counter(t for a in animals for t in a["types"])
    print("\n" + ", ".join(f"{lbl}: {cnt[tid]}" for tid, lbl in Z.TYPES) + f" – zusammen {len(animals)} Tierarten, {len(pairs)} Beziehungen.")
    if errors:
        print(f"\n{len(errors)} Fehler – es wurde nichts geschrieben. Bitte in der Excel korrigieren und erneut starten.")
        sys.exit(1)
    if not animals:
        print("Keine Tierarten eingetragen – die bestehenden Daten bleiben unverändert.")
        return
    if args.pruefen:
        print("Prüfung ohne Fehler (nichts geschrieben).")
        return

    animals = [prune(a) for a in sorted(animals, key=lambda a: a["id"])]
    Z.save_json(meta.get("animalsFile", "data/tiere.json"), {"animals": animals})
    Z.save_json(meta.get("compatFile", "data/compat.json"), {
        "_comment": "Wird aus der Excel (Spalten „Zusammenleben“) erzeugt. benefit: a profitiert von b; neutral: verträglich ohne Bonus. Nicht aufgeführte Paare können nicht zusammenleben.",
        "pairs": pairs})
    meta["dlcs"] = dlcs
    version = read_version(wb)
    if version:
        meta["version"] = version
    updated = read_updated(wb)
    if updated:
        meta["updated"] = updated
    else:
        meta.pop("updated", None)
    counter = read_info(wb, "GoatCounter-Code")
    if counter:  # auch ganzer Einbau-Code oder Link erlaubt, z. B. https://ninipisi.goatcounter.com/count -> ninipisi
        m = re.search(r"([a-z0-9-]+)\.goatcounter\.com", counter.lower())
        if m:
            counter = m.group(1)
    if counter and not re.fullmatch(r"[a-z0-9-]+", counter.strip().lower()):
        print(f"Hinweis  Info: GoatCounter-Code „{counter}“ ungültig (nur Kleinbuchstaben, Ziffern, Bindestrich) – Zähler bleibt aus")
        counter = None
    if counter:
        meta["counter"] = counter.strip().lower()
    else:
        meta.pop("counter", None)
    Z.save_json("data/meta.json", meta)
    print("Daten geschrieben: data/tiere.json, data/compat.json, data/meta.json")


if __name__ == "__main__":
    main()
