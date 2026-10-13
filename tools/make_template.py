"""Erzeugt die Excel-Vorlage für die Zoopedia-Daten.

Aufruf (im Projektordner):
  python tools/make_template.py                      -> leere Vorlage  daten/Zoopedia_PZ2_Daten.xlsx
  python tools/make_template.py --mit-daten          -> Vorlage mit den aktuellen Daten aus data/*.json
  python tools/make_template.py --ziel pfad.xlsx     -> anderer Speicherort
Benötigt: pip install openpyxl
"""
import argparse
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import zoopedia_schema as Z  # noqa: E402

GROUP_COLORS = {"Allgemein": "0F3B43", "Kontinente": "1F6F7A", "Biome": "2E7D4F", "Flächen": "B7791F",
                "Klima": "A3472E", "Körpermaße": "8A6D3B", "Zaun": "6B4FA0", "Sozialverhalten": "3D6EA8", "Lebenszyklus": "5C6B73"}
HEAD_FILL = PatternFill("solid", fgColor="FFE3A8")
THIN = Side(style="thin", color="C9D3D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MAX_ROWS = 400
TODO_FILL = PatternFill("solid", fgColor="FFF59D")


def get(obj, path):
    for k in path.split("."):
        if obj is None:
            return None
        if isinstance(obj, list):
            i = int(k)
            obj = obj[i] if i < len(obj) else None
        else:
            obj = obj.get(k)
    return obj


def cell_value(animal, path, kind, dlc_names):
    if kind == "x":
        group, key = path.split(".")
        return "X" if key in (animal.get(group) or []) else None
    v = get(animal, path)
    if v is None:
        return None
    if kind == "bool":
        return "Ja" if v else "Nein"
    if kind == "status":
        for sid, abbr, _ in Z.STATUSES:  # _ = ausgeschriebener Name
            if sid == v:
                return f"{abbr} ({_})"
        return v
    if kind == "dlc":
        return dlc_names.get(v, v)
    if kind == "measure":
        return dict(Z.MEASURES).get(v, v)
    if kind == "walk":
        return dict(Z.WALK).get(v, v)
    if kind == "val":
        if isinstance(v, list):
            return f"{v[0]}-{v[1]}" if v[0] != v[1] else v[0]
        if v == "death":
            return "Tod"
        if v == "once":
            return "Nur 1 Paarung"
    return v


def build(target, with_data):
    meta = Z.load_json("data/meta.json")
    wb = Workbook()

    # ---- Anleitung ----
    ws = wb.active
    ws.title = "Anleitung"
    lines = [
        ("Zoopedia Planet Zoo 2 – Datenvorlage", True),
        ("", False),
        ("So pflegst du die Daten:", True),
        ("GELB markierte Zellen sind noch offen: Diese Werte nach dem Release im Spiel nachschauen und eintragen (die Markierung bleibt, einfach überschreiben).", True),
        ("Bereits eingetragen sind alle Angaben aus der offiziellen Zoopedia (Namen, wissenschaftliche Namen, Inhalt, Gefährdung, Systematik, Kontinente, Biome, Körpermaße und – wo veröffentlicht – Zaun-Anforderungen).", False),
        ("", False),
        ("1. Je Bereich gibt es ein Blatt (Gehege, Voliere, Aquarium, Vivarium, Begehbares Vivarium). Eine Zeile = eine Tierart. Eine Tierart darf in mehreren Blättern stehen (z. B. Fische in Aquarium und Vivarium).", False),
        ("2. Neue DLC? Zuerst im Blatt „Inhalte“ eine Zeile ergänzen, dann die Tiere mit diesem Inhalt eintragen.", False),
        ("3. Felder mit Auswahlliste (Ja/Nein, X, Gefährdungslevel, Inhalt, Begehbar) bitte über die Liste befüllen.", False),
        ("4. Flächen: Wert für 1 Tier UND Wert für 2 Tiere eintragen. Die Seite rechnet: Mehrbedarf = 2 Tiere − 1 Tier; Bedarf für n Tiere = Basis + (n−1) × Mehrbedarf.", False),
        ("5. Spannen als „von-bis“ schreiben, z. B. 2-3. Sterilität: Zahl, Spanne, „Tod“ oder „Unbekannt“.", False),
        ("6. Die Spalte „ID“ leer lassen – sie entsteht automatisch aus dem Namen.", False),
        ("7. Überschriften (Zeile 2) nicht umbenennen. Eigene Zusatzspalten rechts daneben sind erlaubt und werden ignoriert.", False),
        ("8. Zeilen, deren Tierart mit „Beispiel“ beginnt, werden nicht übernommen.", False),
        ("9. Einheiten der Lebenszyklus-Werte im Blatt „Einheiten“ festlegen (gilt für alle Tiere).", False),
        ("9b. Übersetzte Tiernamen im Blatt „Tiernamen“ (eine Zeile je Art, deutscher Name als Schlüssel). Es ist mit allen bisher bekannten Namen vorbefüllt; leere Felder zeigen den deutschen Namen.", False),
        ("10. Vergesellschaftung im gleichnamigen Blatt, eine Zeile je Beziehung: Tierart | profitiert von / profitieren voneinander / nicht verträglich mit | Partner-Tierart. "
         "„profitiert von“ gilt nur in eine Richtung (z. B. Giraffe profitiert von Zebra, dem Zebra ist die Giraffe egal). Beide Namen über die Auswahlliste wählen.", False),
        ("", False),
        ("Umwandeln in die Webseiten-Daten:", True),
        ("Datei im Projekt unter daten/Zoopedia_PZ2_Daten.xlsx speichern und hochladen – die GitHub Action erzeugt die Daten automatisch.", False),
        ("Lokal geht es auch: python tools/excel_to_json.py daten/Zoopedia_PZ2_Daten.xlsx", False),
        ("Das Skript meldet fehlende oder unplausible Werte mit Blatt und Zeilennummer.", False),
    ]
    for i, (txt, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = Font(bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    # ---- Inhalte ----
    wi = wb.create_sheet("Inhalte")
    heads = ["ID", "Name (Deutsch)", "Name (Englisch)", "Standardmäßig aktiv", "Steam-Link"]
    for c, hname in enumerate(heads, 1):
        cell = wi.cell(row=1, column=c, value=hname)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
    builtin = {"base": ("Basisspiel", "Base game"), "deluxe": ("Deluxe-Edition", "Deluxe Edition")}
    dlc_names = {}
    for r, d in enumerate(meta["dlcs"], 2):
        de, en = builtin.get(d["id"], (None, None))
        names = d.get("names") or {}
        de = names.get("de", de) or d["id"]
        en = names.get("en", en) or de
        dlc_names[d["id"]] = de
        for c, v in enumerate([d["id"], de, en, "Ja" if d.get("default") else "Nein", d.get("url")], 1):
            wi.cell(row=r, column=c, value=v).border = BORDER
    wi.cell(row=1, column=7, value="Hinweis: Eine Tierart gehört genau zu einem Inhalt. „Standardmäßig aktiv“ = beim ersten Besuch angehakt. ID leer lassen = automatisch. "
          "Steam-Link = Shopseite des Inhalts (erscheint in „Meine Spielinhalte“; leer = kein Link).").font = Font(italic=True, color="666666")
    for col, w in zip("ABCDE", (16, 30, 30, 22, 60)):
        wi.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"Ja,Nein"', allow_blank=True)
    wi.add_data_validation(dv)
    dv.add("D2:D200")
    wi.freeze_panes = "A2"

    # ---- Listen (versteckt, Quelle der Auswahllisten) ----
    wl = wb.create_sheet("Listen")
    for r, (_, abbr, name) in enumerate(Z.STATUSES, 1):
        wl.cell(row=r, column=1, value=f"{abbr} ({name})")
    for r, (_, name) in enumerate(Z.WALK, 1):
        wl.cell(row=r, column=2, value=name)
    for r, (_, name) in enumerate(Z.UNIT_NAMES, 1):
        wl.cell(row=r, column=3, value=name)
    for r, (_, name) in enumerate(Z.MEASURES, 1):
        wl.cell(row=r, column=4, value=name)
    wl.sheet_state = "hidden"

    # ---- Tierblätter ----
    for sheet, cat in Z.SHEETS:
        cols = Z.COLUMNS[cat]
        wsx = wb.create_sheet(sheet)
        # Zeile 1: Gruppen (verbunden), Zeile 2: Spalten
        start = 1
        while start <= len(cols):
            grp = cols[start - 1][1]
            end = start
            while end < len(cols) and cols[end][1] == grp:
                end += 1
            cell = wsx.cell(row=1, column=start, value=grp)
            fill = PatternFill("solid", fgColor=GROUP_COLORS.get(grp, "0F3B43"))
            for c in range(start, end + 1):
                wsx.cell(row=1, column=c).fill = fill
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")
            if end > start:
                wsx.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            start = end + 1
        for c, (head, grp, path, kind, hint) in enumerate(cols, 1):
            cell = wsx.cell(row=2, column=c, value=head)
            cell.font = Font(bold=True)
            cell.fill = HEAD_FILL
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            if hint:
                cell.comment = Comment(hint, "Zoopedia")
            letter = get_column_letter(c)
            width = 30 if kind == "text" else 16 if kind in ("id", "dlc", "status") else 7 if kind == "x" else 13
            wsx.column_dimensions[letter].width = width
            rng = f"{letter}3:{letter}{MAX_ROWS}"
            dvx = None
            if kind == "bool":
                dvx = DataValidation(type="list", formula1='"Ja,Nein"', allow_blank=True)
            elif kind == "x":
                dvx = DataValidation(type="list", formula1='"X"', allow_blank=True)
            elif kind == "status":
                dvx = DataValidation(type="list", formula1=f"Listen!$A$1:$A${len(Z.STATUSES)}", allow_blank=True)
            elif kind == "walk":
                dvx = DataValidation(type="list", formula1=f"Listen!$B$1:$B${len(Z.WALK)}", allow_blank=True)
            elif kind == "measure":
                dvx = DataValidation(type="list", formula1=f"Listen!$D$1:$D${len(Z.MEASURES)}", allow_blank=True)
            elif kind == "dlc":
                dvx = DataValidation(type="list", formula1="Inhalte!$B$2:$B$200", allow_blank=True)
            elif kind == "int16":
                dvx = DataValidation(type="whole", operator="between", formula1="1", formula2="6", allow_blank=True,
                                     error="Bitte eine ganze Zahl von 1 bis 6 eingeben.", showErrorMessage=True)
            elif kind == "num":
                dvx = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                     error="Bitte eine Zahl ≥ 0 eingeben.", showErrorMessage=True)
            if dvx is not None:
                wsx.add_data_validation(dvx)
                dvx.add(rng)
            if kind == "id":
                for r in range(3, MAX_ROWS + 1):
                    wsx.cell(row=r, column=c).font = Font(color="888888")
        wsx.row_dimensions[2].height = 48
        wsx.freeze_panes = "B3"

        rows = []
        if with_data:
            try:
                data = Z.load_json(f"data/{_file_of(meta, cat)}")
                rows = sorted(data.get("animals", []), key=lambda a: a["name"].lower())
            except FileNotFoundError:
                rows = []
        elif cat == "enclosure":
            rows = [_example()]
        for r, animal in enumerate(rows, 3):
            for c, (head, grp, path, kind, hint) in enumerate(cols, 1):
                v = cell_value(animal, path, kind, dlc_names)
                if v is not None:
                    wsx.cell(row=r, column=c, value=v)
                elif with_data and grp in Z.RELEASE_GROUPS:
                    wsx.cell(row=r, column=c).fill = TODO_FILL  # noch offen: nach dem Release eintragen

    # ---- Einheiten ----
    we = wb.create_sheet("Einheiten")
    for c, hname in enumerate(["Wert", "Einheit"], 1):
        cell = we.cell(row=1, column=c, value=hname)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
    unit_de = dict(Z.UNIT_NAMES)
    for r, (key, name) in enumerate(Z.UNITS, 2):
        we.cell(row=r, column=1, value=name)
        we.cell(row=r, column=2, value=unit_de.get(meta.get("units", {}).get(key, "years"), "Jahre"))
    dvu = DataValidation(type="list", formula1=f"Listen!$C$1:$C${len(Z.UNIT_NAMES)}", allow_blank=False)
    we.add_data_validation(dvu)
    dvu.add(f"B2:B{len(Z.UNITS) + 1}")
    we.column_dimensions["A"].width = 30
    we.column_dimensions["B"].width = 14

    # ---- Vergesellschaftung (gerichtet: Tierart profitiert von Partner) ----
    wc = wb.create_sheet("Vergesellschaftung")
    heads = [("Tierart", "Die Art, um die es geht (Auswahlliste)."),
             ("Beziehung", "profitiert von = nur die Tierart links hat einen Vorteil (Gegenrichtung egal, Zusammenleben trotzdem möglich)\n"
                           "profitieren voneinander = beide haben einen Vorteil\nnicht verträglich mit = dürfen nicht zusammen"),
             ("Partner-Tierart", "Die andere Art (Auswahlliste)."),
             ("Bemerkung (optional)", "Freitext, erscheint im Planer unter dem Paar.")]
    for c, (hname, hint) in enumerate(heads, 1):
        cell = wc.cell(row=1, column=c, value=hname)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.comment = Comment(hint, "Zoopedia")
    wc.cell(row=1, column=6, value="So liest sich jede Zeile: „Massai-Giraffe | profitiert von | Steppenzebra“. "
                                    "Profitieren beide, eine Zeile mit „profitieren voneinander“. Ohne Eintrag zeigt der Planer „kein Vorteil hinterlegt“.").font = Font(italic=True, color="666666")
    wc.column_dimensions["F"].width = 90
    all_names = sorted({a["name"] for _, cat in Z.SHEETS for a in _animals(meta, cat)}, key=str.lower)
    for r, nm in enumerate(all_names, 1):
        wl.cell(row=r, column=5, value=nm)
    nrange = f"Listen!$E$1:$E${max(1, len(all_names)) + 200}"   # Platz für spätere Arten
    for col in ("A", "C"):
        dvn = DataValidation(type="list", formula1=nrange, allow_blank=True, showErrorMessage=True,
                             error="Bitte eine Tierart aus der Liste wählen (deutscher Name).")
        wc.add_data_validation(dvn)
        dvn.add(f"{col}2:{col}2000")
    dvr = DataValidation(type="list", formula1='"' + ",".join(lab for _, lab in Z.RELATIONS) + '"', allow_blank=True)
    wc.add_data_validation(dvr)
    dvr.add("B2:B2000")
    for col, w in zip("ABCD", (30, 26, 30, 50)):
        wc.column_dimensions[col].width = w
    wc.freeze_panes = "A2"
    if with_data:
        compat = Z.load_json(meta.get("compatFile", "data/compat.json"))
        names = {a["id"]: a["name"] for _, cat in Z.SHEETS for a in _animals(meta, cat)}
        lab = dict(Z.RELATIONS)
        pairs = compat.get("pairs", [])
        benefit = {(p["a"], p["b"]) for p in pairs if p.get("type") == "benefit"}
        rows, done = [], set()
        for p in pairs:
            a, b, typ = p["a"], p["b"], p.get("type")
            if typ == "benefit":
                if (a, b) in done:
                    continue
                if (b, a) in benefit:
                    rows.append((a, lab["mutual"], b, p.get("note"))); done |= {(a, b), (b, a)}
                else:
                    rows.append((a, lab["benefit"], b, p.get("note"))); done.add((a, b))
            elif typ == "bonus":
                rows.append((a, lab["mutual"], b, p.get("note")))
            elif typ == "conflict":
                rows.append((a, lab["conflict"], b, p.get("note")))
        for r, (a, rel, b, note) in enumerate(rows, 2):
            for c, v in enumerate((names.get(a, a), rel, names.get(b, b), note), 1):
                if v:
                    wc.cell(row=r, column=c, value=v)

    # ---- Tiernamen (Übersetzungen) ----
    # Wird immer mit allen bekannten Übersetzungen aus data/*.json vorbefüllt.
    wn = wb.create_sheet(Z.NAMES_SHEET)
    heads = [Z.NAMES_KEY_HEAD] + [Z.name_head(c, l) for c, l in Z.NAME_LANGS]
    for c, hname in enumerate(heads, 1):
        cell = wn.cell(row=1, column=c, value=hname)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        wn.column_dimensions[get_column_letter(c)].width = 30 if c == 1 else 24
    wn.cell(row=1, column=1).comment = Comment("Deutscher Name genau wie in den Tierblättern. Leere Übersetzungen zeigen auf der Seite den deutschen Namen. Zusätzliche Zeilen (Arten, die es noch nicht gibt) stören nicht.", "Zoopedia")
    known = {}
    for _, cat in Z.SHEETS:
        try:
            for a in Z.load_json(f"data/{_file_of(meta, cat)}").get("animals", []):
                if a.get("names"):
                    known[a["name"]] = a["names"]
        except FileNotFoundError:
            pass
    for r, nm in enumerate(sorted(known, key=str.lower), 2):
        wn.cell(row=r, column=1, value=nm)
        for c, (code, _) in enumerate(Z.NAME_LANGS, 2):
            v = known[nm].get(code)
            if v:
                wn.cell(row=r, column=c, value=v)
    wn.freeze_panes = "B2"
    wn.row_dimensions[1].height = 32

    # Reihenfolge: Anleitung, Inhalte, Tierblätter, Tiernamen, Einheiten, Vergesellschaftung, Listen
    order = ["Anleitung", "Inhalte"] + [s for s, _ in Z.SHEETS] + [Z.NAMES_SHEET, "Einheiten", "Vergesellschaftung", "Listen"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 2
    os.makedirs(os.path.dirname(os.path.abspath(target)), exist_ok=True)
    wb.save(target)
    print("Vorlage gespeichert:", target)


def _animals(meta, cat):
    try:
        return Z.load_json(f"data/{_file_of(meta, cat)}").get("animals", [])
    except FileNotFoundError:
        return []


def _file_of(meta, cat):
    for c in meta["categories"]:
        if c["id"] == cat:
            return os.path.basename(c["file"])
    return cat + ".json"


def _example():
    return {"name": "Beispiel – diese Zeile wird ignoriert", "dlc": "base", "status": "lc", "continents": ["africa"], "biomes": ["grassland"],
            "area1": {"size": 500, "climb": 0, "water": 20, "deep": 0}, "area2": {"size": 600, "climb": 0, "water": 30, "deep": 0},
            "tempMin": 5, "tempMax": 35, "fence": {"climbProof": False, "strength": 3, "height": 2, "moat": {"width": 3, "depth": 1}},
            "social": {"min": 3, "max": 15, "maxMales": 1, "maxFemales": 14, "maleOnly": [2, 6], "femaleOnly": [3, 15]},
            "walkthrough": "no", "lifespan": {"m": 20, "f": 22},
            "lifecycle": {"maturity": [2, 3], "sterility": "death", "offspring": [1, 2], "gestation": 9, "interval": 24}}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Excel-Vorlage für die Zoopedia erzeugen")
    ap.add_argument("--mit-daten", action="store_true", help="aktuelle Daten aus data/*.json eintragen")
    ap.add_argument("--ziel", default=os.path.join(Z.ROOT, "daten", "Zoopedia_PZ2_Daten.xlsx"))
    a = ap.parse_args()
    build(a.ziel, a.mit_daten)
